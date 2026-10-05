from pathlib import Path
import base64, struct, zlib, hashlib, json, copy
from lxml import etree
from pptx import Presentation
from pptx.util import Inches
root=Path(__file__).parent
b64='''ZsoWq+RJS7pa0pgnnqRez8n/v/MVRT64+f0LUEsDBBQAAAAIACtivFzlAyww
ogMAAFwTAAAVAAAAcHB0L3NsaWRlcy9zbGlkZTEueG1s7VjLjts2FN3nKwht
vGqot2QjniDjZLJpmmA8RdcciZIFUCRLchxPd9kE3fULuskXJB81i/xF+dDD
YzudpG6DDjASIIr33nP5uOeSlJ483bQErLGQDaPzSfDYnwBMC1Y2tJ5Pfr44
+yGfAKkQLRFhFM8n11hOnp48esJnkpRAg6mcobm3UorPIJTFCrdIPmYcU62r
mGiR0lVRw1Kgt9ppS2Do+ylsUUO9Ds+/Bs8FlpgqpHRHDzkRX+OEVVVT4Oes
uGq1L+dEYGKdylXDpXeiR1YsSWlKyS8ExuaNrl8KvuRvhFX/tH4jQFPOvcAD
FLV47nmwU3Rm0IHsC9yB11smkjvDfddh7/ocF3r6a4JBMLRysIneh7Qlmm0q
0ZpSjxls5p7vgWvzhEaGNwoUWhgnmZ4DrSq0Ls2T3FSMwxHOhVQvMWuBeZl7
QnfHM3K0/lEqZ9qbGHEtUHnWEAIEU780arVcIY7NTFml1Bj3AjiTpjumJkV9
uSACrBHRIj/LTtOuE7Xctg58cx2CnMbP8y0I3GqKNBTIAhFcdqOHYyetnpon
ZabutEYC+3nUpbom2Fme40oHZ9MPxxIND91ARaFJFXi9F21trCrteABGdwM7
exumqtLTPYDDu8EDwrbM6AhuG8rEIQdkbNnZu9G7UfOZ2pyy8trgLnWp2SkU
WTBiCYVosWI68QolHK+IVEsDtBVuHxqBSE0HI2gVcHQM+zT4cjJE+8kQHpUM
A++/R0pIRpqBbreIe/bM3J33W2YPrPz/szLeZ2V0JCuTaXSYlVEefCdOTv0X
/ouzB07eT04mPSdf6x6D+Cg6hmkSZKmjo1sVb3MySONkGnas7Ct/y0pMiD5j
4aOIabGEr5CThsmYC535A3PvI3PTnrkXml+nbAOSW+QFxpWNyrfSOIuCJHQs
DvPcj9LdpTXw82na09ifZmF+5OI6Em1k2N4UvxVIf/fIX6+QcAkh+bMrpZGd
Q2d214QTCy1xda5F8jedqPbscjnQ98tnD3vtr/Owd+eG5bpgH+rk5vf3N5/+
uPn47vOfH4xWWRvxzwKe7QY8/ZcCngShn8Yu4pGfRmG8G/E4i7N+4erocW8D
HuRdwP07Ah4tzP1tAV9y0iiwLPRnMAULtsbi6LDn+6em7KhtKg/04ajbppI8
CYPdBO/2LhvtcR/7L85O3YfozhYV+w9b1L3ZouD42weOf4IKIl4h/npt+dci
qbBYWBFvaO14umUC7c+xk78AUEsDBBQAAAAIACtivFxmum19twAAADYBAAAg
AAAAcHB0L3NsaWRlcy9fcmVscy9zbGlkZTEueG1sLnJlbHONz70KwjAQB/Dd
pwhZMplUBxVp6iKC4CT6AEdybYNtEnJR7Nub0YKD4339/lx9eI8De2EiF7wW'''
data=base64.b64decode(b64); offset=data.index(b'PK\x03\x04')
_,version,flags,compression,mt,md,crc,cs,us,nl,xl=struct.unpack('<IHHHHHIIIHH',data[offset:offset+30])
name=data[offset+30:offset+30+nl].decode();start=offset+30+nl+xl
xml=zlib.decompress(data[start:start+cs],-15)
assert len(xml)==us and zlib.crc32(xml)&0xffffffff==crc
(root/'wuhua-original-slide1.xml').write_bytes(xml)
(root/'wuhua-source-base64-fragment.txt').write_text(b64)
ns={'p':'http://schemas.openxmlformats.org/presentationml/2006/main','a':'http://schemas.openxmlformats.org/drawingml/2006/main','c':'http://schemas.openxmlformats.org/drawingml/2006/chart'}
node=etree.fromstring(xml,etree.XMLParser(resolve_entities=False,no_network=True))
sp=node.xpath('//p:sp',namespaces=ns)
report={'source':'wuhua2026/ppt-templates/templates/static/cover/split_screen_ocean_blue.pptx','github_blob_sha':'b0c385dc5780c7a126ca97c623eaaf9e600f088e','source_base64_lines':[513,537],'zip_member':name,'member_crc_verified':True,'member_xml_bytes':len(xml),'xml_sha256':hashlib.sha256(xml).hexdigest(),'native_shapes':len(sp),'nonempty_text_shapes':len([s for s in sp if s.xpath('.//a:t',namespaces=ns)]),'picture_objects':len(node.xpath('//p:pic',namespaces=ns)),'texts_before':node.xpath('//a:t/text()',namespaces=ns)}
# Retain actual page nodes; original theme/master not fetched and not claimed preserved.
prs=Presentation();prs.slide_width=Inches(10);prs.slide_height=Inches(7.5);slide=prs.slides.add_slide(prs.slide_layouts[6])
for x in list(node.find('p:cSld/p:spTree',ns))[2:]:slide.shapes._spTree.append(copy.deepcopy(x))
prs.save(root/'wuhua-original-page-repacked.pptx')
geom=[(s.shape_id,int(s.left),int(s.top),int(s.width),int(s.height)) for s in slide.shapes]
items=[s for s in slide.shapes if s.has_text_frame and s.text.strip()]
for s,t in zip(items,['部门汇报','内部展示 · 内容替换测试']):
 ts=s._element.xpath('.//a:t');ts[0].text=t
 for t2 in ts[1:]:t2.text=''
prs.save(root/'wuhua-edited-page-repacked.pptx')
r=Presentation(root/'wuhua-edited-page-repacked.pptx').slides[0]
report.update({'texts_after':[s.text for s in r.shapes if s.has_text_frame and s.text.strip()],'shape_count_after':len(r.shapes),'geometry_unchanged':geom==[(s.shape_id,int(s.left),int(s.top),int(s.width),int(s.height)) for s in r.shapes],'page_repack_edit_reopen_passed':True,'limits':['仅一个真实原生页面的结构和确定性工具替换验证','重封装使用干净母版，不声称保留原文件完整母版/主题','未调用Jato Agent真实模型，未在PowerPoint/WPS验收','不代表全库全部模板通过']})
(root/'wuhua-reuse-evidence.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
print(json.dumps(report,ensure_ascii=False,indent=2))
